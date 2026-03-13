#!/usr/bin/env python3
"""
Script de exportação: converte FATO_Lancamentos e DIM_Indicadores do Excel para JSON.
Use para atualizar o Dashboard Gestão à Vista.

USO:
  python export_dashboard.py [caminho_para_xlsx]

Exemplo:
  python export_dashboard.py Indicadores_Gestao_a_Vista_v11.xlsx

O arquivo 'dashboard_data.json' será gerado na mesma pasta.
Basta carregar esse arquivo no Dashboard (aba "Atualizar Dados").

Para atualização automática via URL, suba o dashboard_data.json em um servidor web
ou pasta compartilhada e configure a URL no dashboard.
"""

import sys
import os
import json
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("Instalando openpyxl...")
    os.system("pip install openpyxl --quiet")
    import openpyxl

def to_serializable(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:  # NaN
            return None
    return v

def read_sheet(ws, max_rows=5000):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else f'col_{i}' for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:max_rows]:
        if any(c is not None for c in row):
            obj = {headers[i]: to_serializable(row[i]) for i in range(min(len(headers), len(row)))}
            result.append(obj)
    return result

def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else 'Indicadores_Gestao_a_Vista_v11.xlsx'
    
    if not os.path.exists(xlsx_path):
        print(f"ERRO: Arquivo não encontrado: {xlsx_path}")
        sys.exit(1)
    
    print(f"Lendo arquivo: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    output = {}
    
    # DIM_Indicadores
    if 'DIM_Indicadores' in wb.sheetnames:
        output['DIM_Indicadores'] = read_sheet(wb['DIM_Indicadores'])
        print(f"  DIM_Indicadores: {len(output['DIM_Indicadores'])} registros")
    else:
        output['DIM_Indicadores'] = []
        print("  AVISO: Aba DIM_Indicadores não encontrada!")
    
    # FATO_Lancamentos
    if 'FATO_Lancamentos' in wb.sheetnames:
        output['FATO_Lancamentos'] = read_sheet(wb['FATO_Lancamentos'])
        print(f"  FATO_Lancamentos: {len(output['FATO_Lancamentos'])} lançamentos")
    else:
        output['FATO_Lancamentos'] = []
        print("  AVISO: Aba FATO_Lancamentos não encontrada!")
    
    # DIM_Areas
    if 'DIM_Areas' in wb.sheetnames:
        output['DIM_Areas'] = read_sheet(wb['DIM_Areas'])
        print(f"  DIM_Areas: {len(output['DIM_Areas'])} áreas")
    else:
        output['DIM_Areas'] = []
    
    # Resumo_Texto_Mensal (opcional)
    if 'Resumo_Texto_Mensal' in wb.sheetnames:
        output['Resumo_Texto_Mensal'] = read_sheet(wb['Resumo_Texto_Mensal'])
    
    output['last_updated'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    # Salvar JSON
    out_path = os.path.join(os.path.dirname(xlsx_path), 'dashboard_data.json')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2, default=str)
    
    print(f"\n✅ JSON gerado com sucesso: {out_path}")
    print(f"   Tamanho: {os.path.getsize(out_path)/1024:.1f} KB")
    print(f"\n📋 Próximos passos:")
    print(f"   1. Abra o dashboard.html no navegador")
    print(f"   2. Vá na aba 'Atualizar Dados'")
    print(f"   3. Selecione o arquivo: {out_path}")
    print(f"   OU: Configure uma URL de servidor para atualização automática")

if __name__ == '__main__':
    main()
